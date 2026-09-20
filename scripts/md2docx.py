# -*- coding: utf-8 -*-
"""4단계 — 마크다운 보고서를 Word(.docx) 로 변환한다.

제목·목록·표·코드블록·인용문·그림·구분선, 그리고 문서 안 목차 링크(앵커)와
문서 사이 링크(`[이름](다른파일.md)` → `다른파일.docx`)까지 옮긴다.

사용: python md2docx.py <입력.md> <출력.docx>
여러 개를 한 번에 처리하려면 batch_docx.py 를 쓴다.
"""
import os, re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import kit

_C = kit.load()['docx']
KO = _C['font']
MONO = _C['mono_font']
BODY = _C['body_size']
MAX_IMG_W = _C['max_image_width_in']
PAGE_BREAK = _C['page_break_per_section']


# ---------- low level helpers ----------
def set_font(run, size=None, bold=None, italic=None, color=None, mono=False):
    fname = MONO if mono else KO
    run.font.name = fname
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rf.set(qn(a), fname)
    rf.set(qn('w:eastAsia'), KO)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def shade(pr, fill):
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), fill)
    pr.append(sh)


def shade_par(par, fill):
    shade(par._p.get_or_add_pPr(), fill)


def border(par, edges=('left',), sz=18, color='C9D4E0'):
    pPr = par._p.get_or_add_pPr()
    bd = pPr.find(qn('w:pBdr'))
    if bd is None:
        bd = OxmlElement('w:pBdr')
        pPr.append(bd)
    for e in edges:
        el = OxmlElement('w:' + e)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '6')
        el.set(qn('w:color'), color)
        bd.append(el)


_bm_id = [1000]


def bookmark(par, name):
    _bm_id[0] += 1
    s = OxmlElement('w:bookmarkStart')
    s.set(qn('w:id'), str(_bm_id[0]))
    s.set(qn('w:name'), name)
    e = OxmlElement('w:bookmarkEnd')
    e.set(qn('w:id'), str(_bm_id[0]))
    par._p.insert(0, s)
    par._p.append(e)


def add_external_link(par, target, text):
    """Relative link to a sibling file (md targets are remapped to .docx)."""
    part = par.part
    rid = part.relate_to(target, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), rid)
    _fill_link_run(h, text)
    par._p.append(h)


def add_internal_link(par, anchor, text):
    h = OxmlElement('w:hyperlink')
    h.set(qn('w:anchor'), anchor)
    _fill_link_run(h, text)
    par._p.append(h)


def _fill_link_run(h, text):
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(a), KO)
    rPr.append(rf)
    c = OxmlElement('w:color')
    c.set(qn('w:val'), '1F5FA8')
    rPr.append(c)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '21')
    rPr.append(sz)
    r.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    r.append(t)
    h.append(r)


# ---------- inline markdown ----------
INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*\n]+?\*)', re.S)


def add_inline(par, text, size=BODY, base_bold=False):
    text = re.sub(r'<a id="[^"]*"></a>', '', text)
    text = re.sub(r'<br\s*/?>', ' ', text)
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            set_font(par.add_run(tok[2:-2]), size=size, bold=True)
        elif tok.startswith('`') and tok.endswith('`') and len(tok) > 2:
            set_font(par.add_run(tok[1:-1]), size=size - 0.5, mono=True,
                     color=(0xA3, 0x1D, 0x4E))
        elif tok.startswith('[') and '](' in tok:
            m = re.match(r'\[(.+?)\]\((.+?)\)', tok, re.S)
            if m and m.group(2).startswith('#'):
                add_internal_link(par, 'anchor_' + m.group(2)[1:], m.group(1))
            elif m and m.group(2).lower().endswith('.md'):
                add_external_link(par, m.group(2)[:-3] + '.docx', m.group(1))
            elif m and re.match(r'(https?:|mailto:)', m.group(2)):
                add_external_link(par, m.group(2), m.group(1))
            elif m:
                set_font(par.add_run(m.group(1)), size=size, color=(0x1F, 0x5F, 0xA8))
            else:
                set_font(par.add_run(tok), size=size)
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            set_font(par.add_run(tok[1:-1]), size=size, italic=True)
        else:
            set_font(par.add_run(tok), size=size, bold=base_bold or None)


# ---------- document build ----------
def new_doc():
    d = Document()
    s = d.sections[0]
    s.top_margin = s.bottom_margin = Inches(0.9)
    s.left_margin = s.right_margin = Inches(1.0)
    st = d.styles['Normal']
    st.font.name = KO
    st.font.size = Pt(BODY)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), KO)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.35
    return d


def para(d, space_before=0, space_after=6, indent=0.0):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    return p


def add_inline_heading(par, text, size, color):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            set_font(par.add_run(tok[2:-2]), size=size, bold=True, color=color)
        elif tok.startswith('`') and tok.endswith('`'):
            set_font(par.add_run(tok[1:-1]), size=size - 1, bold=True, mono=True, color=color)
        else:
            set_font(par.add_run(tok), size=size, bold=True, color=color)


def heading(d, level, text, pending_anchor):
    sizes = {1: 18, 2: 14, 3: 11.5}
    colors = {1: (0x10, 0x2A, 0x43), 2: (0x1B, 0x4B, 0x73), 3: (0x2E, 0x2E, 0x2E)}
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt({1: 0, 2: 16, 3: 10}[level])
    p.paragraph_format.space_after = Pt({1: 10, 2: 6, 3: 4}[level])
    p.paragraph_format.keep_with_next = True
    if PAGE_BREAK and level == 2 and len(d.paragraphs) > 1:
        p.paragraph_format.page_break_before = True
    add_inline_heading(p, text, sizes[level], colors[level])
    if level == 2:
        border(p, edges=('bottom',), sz=8, color='B8C6D6')
    if pending_anchor:
        bookmark(p, 'anchor_' + pending_anchor)
    return p


def add_image(d, path, caption):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    if not os.path.exists(path):
        set_font(p.add_run('[그림 없음: %s]' % os.path.basename(path)),
                 size=9, italic=True, color=(0xB0, 0x30, 0x30))
        return False
    with Image.open(path) as im:
        w, _ = im.size
    win = min(MAX_IMG_W, w / 96.0)
    p.add_run().add_picture(path, width=Inches(win))
    if caption:
        c = d.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(10)
        set_font(c.add_run(caption), size=9, italic=True, color=(0x5A, 0x5A, 0x5A))
    return True


def add_code(d, lines):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.line_spacing = 1.1
    shade_par(p, 'F4F6F8')
    border(p, edges=('left',), sz=18, color='9FB3C8')
    for i, ln in enumerate(lines):
        if i:
            p.add_run().add_break()
        set_font(p.add_run(ln), size=9, mono=True, color=(0x24, 0x29, 0x2E))


def add_table(d, rows):
    ncol = max(len(r) for r in rows)
    t = d.add_table(rows=0, cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(ncol):
            cell = cells[ci]
            pf = cell.paragraphs[0].paragraph_format
            pf.space_before = Pt(2)
            pf.space_after = Pt(2)
            pf.line_spacing = 1.2
            txt = row[ci] if ci < len(row) else ''
            add_inline(cell.paragraphs[0], txt, size=9.5, base_bold=(ri == 0))
            if ri == 0:
                shade(cell._tc.get_or_add_tcPr(), 'E8EEF4')
    d.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def hr(d):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    border(p, edges=('bottom',), sz=6, color='D6DEE6')


def convert(md_path, out_path):
    base = os.path.dirname(os.path.abspath(md_path))
    lines = open(md_path, encoding='utf-8').read().split('\n')
    d = new_doc()
    stats = {'headings': 0, 'tables': 0, 'code': 0, 'img_ok': 0, 'img_missing': 0,
             'links': 0, 'anchors': 0}
    i, pending_anchor = 0, None
    while i < len(lines):
        raw = lines[i]
        s = raw.strip()
        m = re.match(r'<a id="([^"]+)"></a>\s*$', s)
        if m:
            pending_anchor = m.group(1)
            stats['anchors'] += 1
            i += 1
            continue
        if not s:
            i += 1
            continue
        if re.match(r'^(---+|\*\*\*+|___+)$', s):
            hr(d)
            i += 1
            continue
        if s.startswith('```'):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                buf.append(lines[i])
                i += 1
            i += 1
            while buf and not buf[-1].strip():
                buf.pop()
            add_code(d, buf if buf else [''])
            stats['code'] += 1
            continue
        m = re.match(r'^(#{1,6})\s+(.*)$', s)
        if m:
            heading(d, min(3, len(m.group(1))), m.group(2).strip(), pending_anchor)
            pending_anchor = None
            stats['headings'] += 1
            i += 1
            continue
        m = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$', s)
        if m:
            p = os.path.normpath(os.path.join(base, m.group(2)))
            ok = add_image(d, p, m.group(1))
            stats['img_ok' if ok else 'img_missing'] += 1
            i += 1
            continue
        if s.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s:\-|]+\|$', lines[i + 1].strip()):
            def split_row(r):
                return [c.strip() for c in r.strip().strip('|').split('|')]
            rows = [split_row(lines[i])]
            i += 2
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(d, rows)
            stats['tables'] += 1
            continue
        if s.startswith('>'):
            buf = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                buf.append(lines[i].strip().lstrip('>').strip())
                i += 1
            p = para(d, space_before=4, space_after=8, indent=0.2)
            shade_par(p, 'FBF8EE')
            border(p, edges=('left',), sz=18, color='D9C48A')
            add_inline(p, ' '.join(buf), size=10)
            continue
        stats['links'] += len(re.findall(r'\]\(#', s))
        m = re.match(r'^(\s*)([-*+])\s+(.*)$', raw)
        if m:
            depth = len(m.group(1)) // 2
            p = para(d, space_after=3, indent=0.25 + 0.22 * depth)
            p.paragraph_format.first_line_indent = Inches(-0.16)
            set_font(p.add_run('• ' if depth == 0 else '– '), size=BODY,
                     color=(0x1B, 0x4B, 0x73))
            add_inline(p, m.group(3), size=BODY)
            i += 1
            continue
        m = re.match(r'^(\s*)(\d+)\.\s+(.*)$', raw)
        if m:
            depth = len(m.group(1)) // 2
            p = para(d, space_after=3, indent=0.25 + 0.22 * depth)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            set_font(p.add_run(m.group(2) + '. '), size=BODY, bold=True,
                     color=(0x1B, 0x4B, 0x73))
            add_inline(p, m.group(3), size=BODY)
            i += 1
            continue
        p = para(d, space_after=6)
        add_inline(p, s, size=BODY)
        i += 1
    d.save(out_path)
    return stats


if __name__ == '__main__':
    src, dst = sys.argv[1], sys.argv[2]
    st = convert(src, dst)
    print(os.path.basename(dst), '|',
          ' '.join('%s=%s' % (k, v) for k, v in st.items()), '|',
          '%.0fKB' % (os.path.getsize(dst) / 1024.0))
