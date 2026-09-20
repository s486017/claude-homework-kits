# -*- coding: utf-8 -*-
"""4단계 검증 — 만들어진 .docx 를 원본 .md 와 대조한다.

확인 항목: 그림 수, 문서 안 목차 링크의 연결 여부, 문서 사이 링크의 대상 존재,
한글 본문 유실. 제출 전에 이것만 통과하면 "그림이 안 보인다" 류 사고는 안 난다.

`--render` 를 주면 Word(윈도우) 나 LibreOffice 로 실제로 열어 PDF 로 내보내
페이지 수와 실제 표시된 그림 수까지 확인한다.

사용: python verify_docx.py [--render]
"""
import argparse
import os
import re
import subprocess
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import kit  # noqa: E402
import docx  # noqa: E402


def check_one(md, dx, dst_dir):
    src = md.read_text(encoding='utf-8')
    z = zipfile.ZipFile(dx)
    x = z.read('word/document.xml').decode('utf-8')
    rels = z.read('word/_rels/document.xml.rels').decode('utf-8')

    shapes = len(re.findall(r'<w:drawing>', x))          # 실제 삽입된 그림
    md_imgs = re.findall(r'!\[[^\]]*\]\(([^)]+)\)', src)
    bms = set(re.findall(r'w:bookmarkStart[^>]*w:name="([^"]+)"', x))
    anchors = re.findall(r'<w:hyperlink w:anchor="([^"]+)"', x)
    md_anchors = re.findall(r'\]\((#[^)]+)\)', src)
    ext = re.findall(r'Type="[^"]*/hyperlink"[^>]*Target="([^"]+)"', rels)
    md_ext = [t for t in re.findall(r'\]\(([^)#][^)]*)\)', src)
              if not t.lower().endswith(('.png', '.jpg', '.jpeg', '.gif'))]

    d = docx.Document(dx)
    text = '\n'.join(p.text for p in d.paragraphs) + '\n' + '\n'.join(
        c.text for t in d.tables for r in t.rows for c in r.cells)
    body = re.sub(r'!\[[^\]]*\]\([^)]*\)', '', src)
    missing_words = sorted({w for w in re.findall(r'[가-힣]{4,}', body) if w not in text})

    problems = []
    if shapes != len(md_imgs):
        problems.append('그림 %d/%d' % (shapes, len(md_imgs)))
    dangling = [a for a in anchors if a not in bms]
    if dangling:
        problems.append('목차 링크 끊김 %s' % dangling[:3])
    if len(anchors) != len(md_anchors):
        problems.append('목차 링크 수 %d/%d' % (len(anchors), len(md_anchors)))
    if len(ext) != len(md_ext):
        problems.append('문서 사이 링크 %d/%d' % (len(ext), len(md_ext)))
    broken = [t for t in ext if not t.startswith(('http', 'mailto'))
              and not (dst_dir / t).exists()]
    if broken:
        problems.append('링크 대상 파일 없음 %s' % broken[:3])
    if missing_words:
        problems.append('본문 누락 %s' % missing_words[:5])
    return {'name': md.stem, 'shapes': shapes, 'anchors': len(anchors),
            'ext': len(ext), 'paras': len(d.paragraphs), 'tables': len(d.tables),
            'problems': problems}


def render_tool():
    """설치돼 있으면 Word(윈도우) 또는 LibreOffice 를 쓴다."""
    if sys.platform == 'win32':
        try:
            import win32com.client  # noqa: F401
            return 'word'
        except ImportError:
            pass
    for exe in ('soffice', '/Applications/LibreOffice.app/Contents/MacOS/soffice',
                r'C:\Program Files\LibreOffice\program\soffice.exe'):
        try:
            subprocess.run([exe, '--version'], capture_output=True, timeout=30)
            return exe
        except (OSError, subprocess.SubprocessError):
            continue
    return None


def render_check(files, out_dir):
    tool = render_tool()
    if not tool:
        print('[건너뜀] Word(pywin32) 나 LibreOffice 가 없어 실제 렌더링 확인은 못 합니다.')
        return
    try:
        import pymupdf
    except ImportError:
        print('[건너뜀] pymupdf 가 없어 렌더링 확인은 못 합니다. (pip install pymupdf)')
        return
    pdf_dir = out_dir / '_pdf'
    pdf_dir.mkdir(exist_ok=True)
    print()
    print('실제 렌더링 확인 (%s)' % ('Word' if tool == 'word' else 'LibreOffice'))
    word = None
    if tool == 'word':
        import win32com.client
        word = win32com.client.Dispatch('Word.Application')
        word.Visible = False
        word.DisplayAlerts = 0
    total_pages = total_imgs = 0
    try:
        for f in files:
            pdf = pdf_dir / (f.stem + '.pdf')
            if word:
                d = word.Documents.Open(str(f.resolve()), ReadOnly=True)
                d.SaveAs(str(pdf), FileFormat=17)
                d.Close(False)
            else:
                subprocess.run([tool, '--headless', '--convert-to', 'pdf',
                                '--outdir', str(pdf_dir), str(f)],
                               capture_output=True, timeout=300)
            doc = pymupdf.open(pdf)
            imgs = sum(len(doc[i].get_images()) for i in range(doc.page_count))
            total_pages += doc.page_count
            total_imgs += imgs
            print('  %-40s %3d쪽  그림 %d장' % (f.stem, doc.page_count, imgs))
            doc.close()
    finally:
        if word:
            word.Quit()
    print('  합계 %d쪽 | 표시된 그림 %d장 | PDF: %s' % (total_pages, total_imgs, pdf_dir))


def main():
    cfg = kit.load()
    ap = argparse.ArgumentParser()
    ap.add_argument('--render', action='store_true', help='Word/LibreOffice 로 실제 열어 확인')
    a = ap.parse_args()

    src = kit.path_for(cfg, 'reports')
    dst = kit.path_for(cfg, 'docx')
    rows, missing = [], []
    files = []
    for md in sorted(src.glob('*.md')):
        dx = dst / (md.stem + '.docx')
        if not dx.exists():
            missing.append(md.stem)
            continue
        files.append(dx)
        rows.append(check_one(md, dx, dst))
    if not rows:
        raise SystemExit('검증할 docx 가 없습니다. 먼저 python batch_docx.py')

    print('%-40s %6s %6s %6s %6s %6s' % ('문서', '그림', '목차링크', '문서간', '문단', '표'))
    for r in rows:
        print('%-40s %6d %6d %6d %6d %6d'
              % (r['name'], r['shapes'], r['anchors'], r['ext'], r['paras'], r['tables']))
    bad = [r for r in rows if r['problems']]
    print()
    print('문서 %d개 | 그림 합계 %d장 | 문제 %d건'
          % (len(rows), sum(r['shapes'] for r in rows), len(bad) + len(missing)))
    for n in missing:
        print('  [문제] %s — docx 가 없음' % n)
    for r in bad:
        print('  [문제] %s — %s' % (r['name'], ', '.join(r['problems'])))

    if a.render:
        render_check(files, dst)
    if not bad and not missing:
        print()
        print('제출 준비 완료: %s' % dst)
    sys.exit(1 if (bad or missing) else 0)


if __name__ == '__main__':
    main()
